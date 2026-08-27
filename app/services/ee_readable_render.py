"""Deterministic Express Estimate JSON → readable markdown / Excel (no AI)."""

from __future__ import annotations

import io
import json
import re
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font


_SECTION_TITLES: dict[str, str] = {
    "projectDetails": "Project details",
    "exterior": "Exterior",
    "foundation": "Foundation",
    "rooms": "Rooms",
}

# Site form field order (camelCase aliases as stored in payload JSON).
_ORDER_PROJECT_DETAILS: tuple[str, ...] = (
    "insuredName",
    "claimNumber",
    "street",
    "city",
    "zipCode",
    "depreciationRange",
    "projectName",
    "inspectionDate",
    "propertyAddress",
    "propertyType",
    "preFirm",
    "adjusterName",
    "notes",
)

_ORDER_EXTERIOR: tuple[str, ...] = (
    "pressureWash",
    "dumpster",
    "hvac",
    "electrical",
    "finishes",
)

_ORDER_EXTERIOR_FINISHES: tuple[str, ...] = (
    "exteriorPaint",
    "siding",
    "sheathing",
    "houseWrap",
    "backerBoard",
    "wallInsulation",
)

_ORDER_EXTERIOR_HVAC: tuple[str, ...] = (
    "condenserUnits",
    "packageUnits",
    "miniSplits",
)

_ORDER_EXTERIOR_ELECTRICAL: tuple[str, ...] = (
    "exteriorOutlets",
    "disconnect30Amp",
    "breakerPanel",
    "meterBox",
    "meterBoxQty",
    "meterBoxSize",
)

_ORDER_FOUNDATION: tuple[str, ...] = (
    "crawlspace",
    "insulation",
    "enclosureRemoval",
    "sumpPump",
    "waterHeater",
    "waterSoftener",
    "subgradeAreaCoverage",
    "hvac",
    "basement",
    "electrical",
    "stairs",
    "elevator",
)

_ORDER_FOUNDATION_HVAC: tuple[str, ...] = (
    "airHandlers",
    "boiler",
    "furnace",
    "baseboardHeat",
)

_ORDER_FOUNDATION_ELECTRICAL: tuple[str, ...] = (
    "outlets110",
    "outlets220",
    "gfiOutlets",
    "lightSwitch",
    "junctionBox",
    "breakerPanel",
    "meterBox",
    "meterBoxQty",
    "meterBoxSize",
    "houseRewire",
)

_ORDER_SUBGRADE: tuple[str, ...] = (
    "drywall",
    "wallInsulation",
    "foundationalDoor",
    "foundationalWindowsEnabled",
    "foundationalWindows",
)

_ORDER_ENCLOSURE_REMOVAL: tuple[str, ...] = (
    "sandRemoval",
    "backfill",
    "confinedSpace",
)

_ORDER_ROOM: tuple[str, ...] = (
    "id",
    "name",
    "type",
    "sqft",
    "nfipCleaning",
    "flooring",
    "trim",
    "wallCovering",
    "windowsEnabled",
    "windows",
    "electrical",
    "vanity",
    "pedestalSink",
    "toilet",
    "shower",
    "cabinets",
    "countertop",
    "plumbing",
    "appliances",
    "doorsEnabled",
    "doors",
    "notes",
)

_ORDER_NFIP_CLEANING: tuple[str, ...] = ("enabled", "wall", "floor")
_ORDER_NFIP_WALL: tuple[str, ...] = ("height", "wallType", "ceilingAffected")
_ORDER_NFIP_FLOOR: tuple[str, ...] = ("type", "areaOnCrawlspace")

_ORDER_FLOORING: tuple[str, ...] = (
    "enabled",
    "multipleLayers",
    "layers",
    "vaporBarrier",
    "subfloorReplacement",
    "f9Note",
)

_ORDER_FLOOR_LAYER: tuple[str, ...] = (
    "id",
    "type",
    "grade",
    "application",
    "action",
    "vaporBarrier",
    "subfloorReplacement",
)

_ORDER_TRIM: tuple[str, ...] = (
    "enabled",
    "baseboardHeight",
    "material",
    "detail",
    "finish",
    "cap",
    "shoe",
    "shoeFinish",
    "subtractCabinetry",
    "vinylCoveEnabled",
    "vinylCoveSize",
    "tileBaseEnabled",
    "tileBaseGrade",
)

_ORDER_WALL_COVERING: tuple[str, ...] = (
    "enabled",
    "material",
    "type",
    "replacementHeight",
    "fullWall",
    "ceilingReplacementAddon",
    "texture",
    "textureType",
    "panelingStyle",
    "panelingFinish",
    "panelingGrade",
    "chairRailAction",
    "chairRailFinish",
)

_ORDER_ROOM_ELECTRICAL: tuple[str, ...] = (
    "enabled",
    "outlets110",
    "outlets220",
    "gfiOutlets",
    "lightSwitches",
    "ceilingLights",
    "ceilingFans",
    "bathroomLightBar",
    "bathroomLightBarQty",
)

_ORDER_WINDOW: tuple[str, ...] = (
    "id",
    "type",
    "material",
    "size",
    "grade",
    "quantity",
    "finish",
    "blinds",
    "casingTrim",
    "marbleSillReplace",
    "marbleSillDetach",
)

_ORDER_DOOR: tuple[str, ...] = (
    "id",
    "category",
    "type",
    "size",
    "grade",
    "finish",
    "handleAction",
    "misc",
    "peepHole",
    "mailSlot",
    "nonCased",
    "casedOpening",
    "casingOpeningSize",
    "casingFinish",
    "sidelites",
    "sidelitesQty",
    "sidelitesSize",
    "sidelitesGrade",
    "panelSize",
    "panelGrade",
    "stormdoorAssembly",
    "retrofitInStucco",
    "cleanSensor",
    "replaceSensor",
)

_ORDER_VANITY: tuple[str, ...] = (
    "enabled",
    "size",
    "grade",
    "custom",
    "detachAndReset",
    "countertop",
    "backsplashUnattached",
    "backsplashAction",
)

_ORDER_VANITY_COUNTERTOP: tuple[str, ...] = (
    "type",
    "grade",
    "size",
    "pStop",
    "sink",
    "action",
    "faucet",
    "faucetAction",
)

_ORDER_KITCHEN_COUNTERTOP: tuple[str, ...] = (
    "enabled",
    "type",
    "grade",
    "size",
    "detachAndReset",
    "action",
    "subdeckReplacement",
)

_ORDER_CABINETS: tuple[str, ...] = (
    "enabled",
    "size",
    "grade",
    "detachAndReset",
    "toeKick",
    "fullHeight",
)

_ORDER_SHOWER: tuple[str, ...] = (
    "enabled",
    "type",
    "detachAndReset",
    "showerFaucet",
    "actionForTub",
    "jetted",
    "jettedMotorReplace",
    "surround",
    "tubShowerFaucet",
    "mortarBedReplace",
    "mortarBedSize",
    "tileCurb",
    "tileCurbSize",
    "walls",
    "tileBench",
    "tileNiche",
    "tileNicheQty",
    "towelBar",
    "tileSoapDish",
    "tileSoapDishQty",
    "grabBar",
    "grabBarQty",
    "tileFeatureStrip",
    "glassDoor",
    "glassDoorAction",
)

_ORDER_APPLIANCES: tuple[str, ...] = (
    "enabled",
    "refrigerator",
    "dishwasher",
    "range",
    "cooktop",
    "waterHeater",
    "wallOven",
    "airHandler",
    "boiler",
    "furnace",
    "baseboardHeat",
)

_ORDER_PLUMBING: tuple[str, ...] = (
    "replaceFaucetSink",
    "drFaucetSink",
    "waterSupplyLine",
    "reverseOsmosis",
    "garbageDisposal",
)

_ORDER_BREAKER_PANEL: tuple[str, ...] = (
    "enabled",
    "amps",
    "arcFaults",
    "panelReplacement",
    "circuitReplacement",
    "panelType",
    "circuits",
)

# Unique last-path-segment → order (no exterior/foundation ambiguity).
_ORDERS_BY_KEY: dict[str, tuple[str, ...]] = {
    "projectDetails": _ORDER_PROJECT_DETAILS,
    "exterior": _ORDER_EXTERIOR,
    "foundation": _ORDER_FOUNDATION,
    "finishes": _ORDER_EXTERIOR_FINISHES,
    "subgradeAreaCoverage": _ORDER_SUBGRADE,
    "enclosureRemoval": _ORDER_ENCLOSURE_REMOVAL,
    "nfipCleaning": _ORDER_NFIP_CLEANING,
    "wall": _ORDER_NFIP_WALL,
    "floor": _ORDER_NFIP_FLOOR,
    "flooring": _ORDER_FLOORING,
    "trim": _ORDER_TRIM,
    "wallCovering": _ORDER_WALL_COVERING,
    "vanity": _ORDER_VANITY,
    "shower": _ORDER_SHOWER,
    "appliances": _ORDER_APPLIANCES,
    "plumbing": _ORDER_PLUMBING,
    "cabinets": _ORDER_CABINETS,
    "breakerPanel": _ORDER_BREAKER_PANEL,
}

# Path suffix → order for ambiguous or room-scoped keys.
_ORDERS_BY_PATH_SUFFIX: dict[tuple[str, ...], tuple[str, ...]] = {
    ("exterior", "hvac"): _ORDER_EXTERIOR_HVAC,
    ("exterior", "electrical"): _ORDER_EXTERIOR_ELECTRICAL,
    ("foundation", "hvac"): _ORDER_FOUNDATION_HVAC,
    ("foundation", "electrical"): _ORDER_FOUNDATION_ELECTRICAL,
    ("rooms",): _ORDER_ROOM,
    ("rooms", "electrical"): _ORDER_ROOM_ELECTRICAL,
    ("rooms", "windows"): _ORDER_WINDOW,
    ("rooms", "doors"): _ORDER_DOOR,
    ("rooms", "vanity", "countertop"): _ORDER_VANITY_COUNTERTOP,
    ("rooms", "countertop"): _ORDER_KITCHEN_COUNTERTOP,
    ("foundation", "subgradeAreaCoverage", "foundationalWindows"): _ORDER_WINDOW,
    ("rooms", "flooring", "layers"): _ORDER_FLOOR_LAYER,
}


def _ordered_keys(obj: dict[str, Any], order: tuple[str, ...] | None) -> list[str]:
    """Emit keys in form order when known; unknown keys keep insertion order after."""
    if not order:
        return list(obj.keys())
    seen: set[str] = set()
    out: list[str] = []
    for k in order:
        if k in obj:
            out.append(k)
            seen.add(k)
    for k in obj:
        if k not in seen:
            out.append(k)
    return out


def _order_for(path: tuple[str, ...]) -> tuple[str, ...] | None:
    if not path:
        return None
    for i in range(len(path)):
        suffix = path[i:]
        if suffix in _ORDERS_BY_PATH_SUFFIX:
            return _ORDERS_BY_PATH_SUFFIX[suffix]
    leaf = path[-1]
    return _ORDERS_BY_KEY.get(leaf)


def _should_skip_branch(val: Any) -> bool:
    """Omit whole objects toggled off on the form (`enabled: false`)."""
    return isinstance(val, dict) and val.get("enabled") is False


def _has_exportable_fields(val: dict[str, Any]) -> bool:
    """True when a list/object row has anything besides a client list-row `id`."""
    for k, v in val.items():
        if k == "id":
            continue
        if _should_skip_branch(v):
            continue
        return True
    return False


def _singularize_words(words: str) -> str:
    """Light plural→singular for list heads: Layers→Layer, Air Handlers→Air Handler."""
    parts = words.split()
    if not parts:
        return words
    last = parts[-1]
    if last.endswith("ies") and len(last) > 3:
        parts[-1] = last[:-3] + "y"
    elif last.endswith("ses") or last.endswith("sses"):
        parts[-1] = last[:-2]
    elif last.endswith("s") and not last.endswith("ss") and len(last) > 1:
        parts[-1] = last[:-1]
    return " ".join(parts)


def _list_entry_value(list_label: str, index: int) -> str:
    """Human value for an id-only list stub, e.g. Layers + 1 → 'Layer 1'."""
    return f"{_singularize_words(list_label)} {index}"


def _indent(level: int) -> str:
    return "  " * level


def _render_value(val: Any, level: int, path: tuple[str, ...] = ()) -> list[str]:
    ind = _indent(level)
    if val is None:
        return [f"{ind}- _(none)_"]
    if isinstance(val, bool):
        return [f"{ind}- `{str(val).lower()}`"]
    if isinstance(val, (int, float)):
        return [f"{ind}- `{val}`"]
    if isinstance(val, str):
        s = val if val else "_(empty)_"
        return [f"{ind}- {s}"]
    if isinstance(val, list):
        if not val:
            return [f"{ind}- _(empty list)_"]
        out: list[str] = []
        for i, item in enumerate(val):
            if isinstance(item, dict):
                if _should_skip_branch(item):
                    continue
                label = item.get("name") or item.get("id")
                title = f"Entry {i + 1}" + (f" — {label}" if label is not None else "")
                out.append(f"{ind}- **{title}**")
                # Room entries and other list objects use the list's path for field order.
                out.extend(_render_value(item, level + 1, path))
            else:
                out.extend(_render_value(item, level, path))
        return out if out else [f"{ind}- _(empty list)_"]
    if isinstance(val, dict):
        if not val:
            return [f"{ind}- _(empty)_"]
        out = []
        for k in _ordered_keys(val, _order_for(path)):
            v = val[k]
            if _should_skip_branch(v):
                continue
            key_label = str(k)
            child_path = path + (k,)
            if isinstance(v, (dict, list)):
                out.append(f"{ind}- **{key_label}**")
                out.extend(_render_value(v, level + 1, child_path))
            elif v is None:
                out.append(f"{ind}- **{key_label}**: _(none)_")
            elif isinstance(v, bool):
                out.append(f"{ind}- **{key_label}**: `{str(v).lower()}`")
            elif isinstance(v, (int, float)):
                out.append(f"{ind}- **{key_label}**: `{v}`")
            elif isinstance(v, str):
                out.append(f"{ind}- **{key_label}**: {v if v else '_(empty)_'}")
            else:
                out.append(f"{ind}- **{key_label}**: `{json.dumps(v, default=str)}`")
        return out if out else [f"{ind}- _(empty)_"]
    return [f"{ind}- `{json.dumps(val, default=str)}`"]


def render_payload_markdown(payload: dict[str, Any]) -> str:
    """Build human-readable markdown from stored EE payload."""
    title = "Express Estimate"
    pd = payload.get("projectDetails")
    if isinstance(pd, dict):
        for key in ("projectName", "insuredName", "claimNumber"):
            val = pd.get(key)
            if isinstance(val, str) and val.strip():
                title = f"Express Estimate — {val.strip()}"
                break

    lines: list[str] = [f"# {title}", ""]
    known_order = ("projectDetails", "exterior", "foundation", "rooms")
    seen: set[str] = set()

    for key in known_order:
        if key not in payload:
            continue
        seen.add(key)
        section = _SECTION_TITLES.get(key, key)
        lines.append(f"## {section}")
        lines.append("")
        chunk = payload[key]
        if chunk is None:
            lines.append("_Not provided._")
        elif _should_skip_branch(chunk):
            lines.append("_Not provided._")
        else:
            lines.extend(_render_value(chunk, 0, (key,)))
        lines.append("")

    other_keys = [k for k in payload if k not in seen]
    for key in other_keys:
        lines.append(f"## {key}")
        lines.append("")
        chunk = payload[key]
        if chunk is None:
            lines.append("_Not provided._")
        elif isinstance(chunk, (dict, list)):
            lines.extend(_render_value(chunk, 0, (key,)))
        else:
            lines.append(json.dumps(chunk, indent=2, default=str))
        lines.append("")

    return "\n".join(lines).rstrip() + "\n"


_JOIN = " » "


def _humanize_key(key: str) -> str:
    """camelCase / snake_case-ish → Title Words for spreadsheet Field column."""
    s = str(key).strip()
    if not s:
        return s
    if s.lower() == "id":
        return "ID"
    s = re.sub(r"[_-]+", " ", s)
    s = re.sub(r"([a-z0-9])([A-Z])", r"\1 \2", s)
    s = re.sub(r"([A-Z]+)([A-Z][a-z])", r"\1 \2", s)
    parts = [p for p in s.split() if p]
    if not parts:
        return str(key)
    out: list[str] = []
    for i, p in enumerate(parts):
        if p.isupper() and len(p) <= 4:
            out.append(p)
        elif i == 0:
            out.append(p[0].upper() + p[1:].lower() if len(p) > 1 else p.upper())
        else:
            out.append(p.lower() if p.islower() or p.isupper() else p[0].upper() + p[1:])
    return " ".join(out)


def _cell_value(val: Any) -> str:
    if val is None:
        return "—"
    if isinstance(val, bool):
        return "Yes" if val else "No"
    if isinstance(val, float):
        if val.is_integer():
            return str(int(val))
        return str(val)
    if isinstance(val, (int, str)):
        return str(val)
    return json.dumps(val, default=str)


def _section_title(key: str) -> str:
    return _SECTION_TITLES.get(key, _humanize_key(key))


def _walk_section(
    section: str,
    val: Any,
    label_parts: list[str],
    rows: list[tuple[str, str, str]],
    path: tuple[str, ...] = (),
) -> None:
    if val is None:
        rows.append((section, _JOIN.join(label_parts) if label_parts else "(empty)", "—"))
        return
    if _should_skip_branch(val):
        return
    if isinstance(val, (bool, int, float, str)):
        rows.append((section, _JOIN.join(label_parts) if label_parts else "(value)", _cell_value(val)))
        return
    if isinstance(val, dict):
        if not val:
            rows.append((section, _JOIN.join(label_parts) if label_parts else "(empty)", "—"))
            return
        for k in _ordered_keys(val, _order_for(path)):
            if k == "id":
                continue
            child = val[k]
            if _should_skip_branch(child):
                continue
            _walk_section(section, child, label_parts + [_humanize_key(k)], rows, path + (k,))
        return
    if isinstance(val, list):
        if not val:
            rows.append((section, _JOIN.join(label_parts) if label_parts else "(list)", "—"))
            return
        for i, item in enumerate(val):
            if isinstance(item, dict):
                if _should_skip_branch(item):
                    continue
                nm = item.get("name")
                extra = ""
                if isinstance(nm, str) and nm.strip():
                    extra = f" — {nm.strip()}"
                # Only the top-level rooms array uses "Room N"; nested lists
                # (windows, doors, layers, …) keep their own list-key label.
                if section == "Rooms" and not label_parts:
                    head = f"Room {i + 1}{extra}"
                    list_label = "Room"
                elif label_parts:
                    list_label = label_parts[-1]
                    head = f"{list_label} {i + 1}{extra}"
                else:
                    list_label = "Entry"
                    head = f"Entry {i + 1}{extra}"
                base = label_parts[:-1] + [head] if label_parts else [head]
                # Id-only stubs ({ "id": … }) mean the user added this list entry
                # with no other fields — show "Layer 1" / "Condenser Unit 1", not the id.
                if not _has_exportable_fields(item):
                    rows.append(
                        (
                            section,
                            _JOIN.join(base),
                            _list_entry_value(list_label, i + 1),
                        )
                    )
                    continue
                _walk_section(section, item, base, rows, path)
            else:
                ip = (label_parts + [f"Item {i + 1}"]) if label_parts else [f"Item {i + 1}"]
                _walk_section(section, item, ip, rows, path)
        return
    rows.append((section, _JOIN.join(label_parts) if label_parts else "(value)", _cell_value(val)))


def _user_friendly_rows(payload: dict[str, Any]) -> list[tuple[str, str, str]]:
    """Section, Field (humanized), Value — one block per top-level key like markdown."""
    rows: list[tuple[str, str, str]] = []
    known_order = ("projectDetails", "exterior", "foundation", "rooms")
    seen: set[str] = set()
    for key in known_order:
        if key not in payload:
            continue
        seen.add(key)
        sec = _section_title(key)
        chunk = payload[key]
        if chunk is None:
            rows.append((sec, "—", "Not provided"))
        else:
            _walk_section(sec, chunk, [], rows, (key,))
    for key in (k for k in payload if k not in seen):
        sec = _section_title(key)
        chunk = payload[key]
        if chunk is None:
            rows.append((sec, "—", "Not provided"))
        elif isinstance(chunk, (dict, list)):
            _walk_section(sec, chunk, [], rows, (key,))
        else:
            rows.append((sec, _humanize_key(key), _cell_value(chunk)))
    return rows


def render_payload_xlsx_bytes(payload: dict[str, Any]) -> bytes:
    """Workbook: title row + Section | Field | Value (humanized, grouped)."""
    title = "Express Estimate"
    pd = payload.get("projectDetails")
    if isinstance(pd, dict):
        for key in ("projectName", "insuredName", "claimNumber"):
            val = pd.get(key)
            if isinstance(val, str) and val.strip():
                title = f"Express Estimate — {val.strip()}"
                break

    data_rows = _user_friendly_rows(payload)

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.merge_cells("A1:C1")
    top = ws["A1"]
    top.value = title
    top.font = Font(bold=True, size=14)
    top.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 28
    ws.append([])
    ws.append(["Section", "Field", "Value"])
    for c in range(1, 4):
        cell = ws.cell(row=3, column=c)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    for section, field, value in data_rows:
        ws.append([section, field, value])
    for row in ws.iter_rows(min_row=4, max_row=ws.max_row, min_col=1, max_col=3):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 48
    ws.column_dimensions["C"].width = 40
    ws.freeze_panes = "A4"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
